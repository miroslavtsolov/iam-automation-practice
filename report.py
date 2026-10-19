from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reconcile import reconcile
from logger import get_logger

log = get_logger()

def build_report():
    merged, orphaned, terminated_but_enabled = reconcile()

    wb = Workbook()
    ws = wb.active
    ws.title = "Access Review"

    headers = ["username", "email", "enabled", "full_name", "status", "_merge"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for _, row in merged.iterrows():
        ws.append([row.get(h) for h in headers])
        if row["_merge"] == "left_only" or (row.get("status") == "terminated" and row.get("enabled") is True):
            for cell in ws[ws.max_row]:
                cell.fill = red_fill

    wb.save("access_review_report.xlsx")
    log.info("Report saved to access_review_report.xlsx")

if __name__ == "__main__":
    build_report()
